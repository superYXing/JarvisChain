from langchain.tools import BaseTool
from pydantic import BaseModel, Field
import json
from src.utils.logger import get_logger
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
import os
from datetime import datetime

# Get logger
logger = get_logger('ppt_add')



class PPTAddTool(BaseTool):
    name: str = "ppt_add_tool"
    description: str = """用于创建和添加PPT元素的工具，包括新建幻灯片、添加文本、图片等。"""
    ppts_dir: str = os.path.join(os.getcwd(), "ppts")

    def __init__(self, **data):
        super().__init__(**data)
        if not os.path.exists(self.ppts_dir):
            os.makedirs(self.ppts_dir)

    def _generate_filename(self, prefix: str = "presentation") -> str:
        """生成唯一的文件名"""
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        return f"{prefix}_{timestamp}.pptx"

    def _create_new_presentation(self, filename: str = None) -> tuple[Presentation, str]:
        """创建新的演示文稿"""
        if not filename:
            filename = self._generate_filename()
        filepath = os.path.join(self.ppts_dir, filename)
        prs = Presentation()
        return prs, filename

    def _save_presentation(self, prs: Presentation, filename: str) -> str:
        """保存演示文稿"""
        filepath = os.path.join(self.ppts_dir, filename)
        prs.save(filepath)
        return filepath

    def _create_presentation(self, filename: str = None) -> tuple[Presentation, str]:
        """创建或打开演示文稿"""
        if not filename:
            filename = self._generate_filename()
        filepath = os.path.join(self.ppts_dir, filename)
        if os.path.exists(filepath):
            return Presentation(filepath), filename
        return self._create_new_presentation(filename)

    def _add_slide(self, prs: Presentation, layout_name: str = "标题和内容") -> None:
        """添加新幻灯片"""
        layouts = {layout.name: layout for layout in prs.slide_layouts}
        layout = layouts.get(layout_name, prs.slide_layouts[1])
        prs.slides.add_slide(layout)

    def _add_text(self, slide, text: str, position: dict, style: dict) -> None:
        """添加文本"""
        left = Inches(position.get('left', 1))
        top = Inches(position.get('top', 1))
        width = Inches(position.get('width', 8))
        height = Inches(position.get('height', 2))

        textbox = slide.shapes.add_textbox(left, top, width, height)
        text_frame = textbox.text_frame
        text_frame.text = text

        # 设置文本样式
        paragraph = text_frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.CENTER if style.get('align') == 'center' else PP_ALIGN.LEFT
        
        # 设置字体
        font = paragraph.font
        font.name = style.get('font', '微软雅黑')
        font.size = Pt(style.get('size', 18))
        
        # 设置颜色
        if 'color' in style:
            r, g, b = style['color']
            font.color.rgb = RGBColor(r, g, b)

    def _add_image(self, slide, image_path: str, position: dict) -> None:
        """添加图片"""
        left = Inches(position.get('left', 1))
        top = Inches(position.get('top', 1))
        width = Inches(position.get('width', 6))
        height = Inches(position.get('height', 4))
        
        slide.shapes.add_picture(image_path, left, top, width, height)

    def _set_background(self, slide, color: tuple) -> None:
        """设置背景颜色"""
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(*color)

    def _run(self, input_data: str) -> str:
        """处理添加PPT元素的命令"""
        try:
            from src.models.llm_models import INTENT_MODEL
            from langchain.schema import SystemMessage, HumanMessage

            system_prompt = """你是一个PPT助手，负责将自然语言命令转换为结构化的JSON操作。
                支持的操作：
                - create: 创建新演示文稿
                - add_slide: 添加幻灯片
                - add_text: 添加文本
                - add_image: 添加图片
                - set_background: 设置背景
                
                返回JSON格式：
                {
                    "operation": "操作名称",
                    "parameters": {
                        "filename": "文件名",
                        "element_type": "元素类型",
                        "content": "内容",
                        "position": {"left": 1, "top": 1, "width": 8, "height": 2},
                        "style": {
                            "font": "字体",
                            "size": 18,
                            "color": [255, 0, 0],
                            "align": "center"
                        }
                    }
                }"""

            messages = [
                SystemMessage(content=system_prompt),
                HumanMessage(content=input_data)
            ]

            response = INTENT_MODEL.invoke(messages)
            operation_json = json.loads(response.content)
            
            operation = operation_json.get("operation")
            parameters = operation_json.get("parameters", {})
            
            # 执行操作
            if operation == "create":
                prs, filename = self._create_new_presentation(parameters.get("filename"))
            else:
                prs, filename = self._create_presentation(parameters.get("filename"))
            
            if operation == "add_slide":
                self._add_slide(prs, parameters.get("layout", "标题和内容"))
            elif operation == "add_text":
                slide = prs.slides[-1]
                self._add_text(slide, parameters["content"], 
                             parameters.get("position", {}),
                             parameters.get("style", {}))
            elif operation == "add_image":
                slide = prs.slides[-1]
                self._add_image(slide, parameters["image_path"],
                              parameters.get("position", {}))
            elif operation == "set_background":
                slide = prs.slides[-1]
                self._set_background(slide, parameters["color"])
            
            # 保存演示文稿
            filepath = self._save_presentation(prs, filename)
            
            return json.dumps({
                "success": True,
                "operation": operation,
                "message": f"成功执行{operation}操作",
                "filepath": filepath
            })

        except Exception as e:
            logger.error(f"PPT添加工具错误: {str(e)}")
            return json.dumps({
                "success": False,
                "message": f"错误: {str(e)}"
            })

    async def _arun(self, input_data: str) -> str:
        return self._run(input_data) 