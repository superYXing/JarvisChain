import json
import os
from pptx import Presentation

input_data = input("请输入:")
data = json.loads(input_data)
filename = data.get("filename")
filename = filename + ".pptx"
layout_name = data.get("layout", "标题和内容")

filepath = "E:\JarvisChain\ppts\presentation_20250527_111101.pptx"
prs = Presentation(filepath)

layout = 1
slide = prs.slides.add_slide(layout)

prs.save(filepath)