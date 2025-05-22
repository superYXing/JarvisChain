from langchain.tools import BaseTool
from pydantic import BaseModel, Field
import json
from src.JarvisChain.utils.logger import get_logger
from pptx import Presentation
import os

# Get logger
logger = get_logger('ppt_remove')



class PPTRemoveTool(BaseTool):
    name: str = "ppt_remove_tool"
    description: str = """用于删除PPT元素的工具，包括删除幻灯片、文本、图片等。"""

    ppts_dir: str = os.path.join(os.getcwd(), "ppts")

    def __init__(self, **data):
        super().__init__(**data)
        if not os.path.exists(self.ppts_dir):
            os.makedirs(self.ppts_dir)

    def _load_presentation(self, filename: str) -> Presentation:
        """加载演示文稿"""
        filepath = os.path.join(self.ppts_dir, filename)
        if not os.path.exists(filepath):
            raise FileNotFoundError(f"未找到演示文稿: {filename}")
        return Presentation(filepath)

    def _save_presentation(self, prs: Presentation, filename: str) -> str:
        """保存演示文稿"""
        filepath = os.path.join(self.ppts_dir, filename)
        prs.save(filepath)
        return filepath

    def _remove_slide(self, prs: Presentation, slide_index: int) -> None:
        """删除指定幻灯片"""
        if 0 <= slide_index < len(prs.slides):
            xml_slides = prs.slides._sldIdLst
            xml_slides.remove(xml_slides[slide_index])

    def _remove_shape(self, slide, shape_index: int) -> None:
        """删除指定形状"""
        if 0 <= shape_index < len(slide.shapes):
            sp = slide.shapes[shape_index]._element
            sp.getparent().remove(sp)

    def _run(self, input_data: str) -> str:
        """处理删除PPT元素的命令"""
        try:
            # 解析 JSON 查询
            data = json.loads(input_data)
            operation = data.get("operation")
            parameters = data.get("parameters", {})
            
            # 执行操作
            prs = self._load_presentation(parameters.get("filename"))
            
            if operation == "remove_slide":
                slide_index = parameters.get("slide_index", 0)
                self._remove_slide(prs, slide_index)
            elif operation == "remove_shape":
                slide_index = parameters.get("slide_index", 0)
                shape_index = parameters.get("shape_index", 0)
                if 0 <= slide_index < len(prs.slides):
                    self._remove_shape(prs.slides[slide_index], shape_index)
            
            # 保存演示文稿
            filepath = self._save_presentation(prs, parameters.get("filename"))
            
            return json.dumps({
                "success": True,
                "operation": operation,
                "message": f"成功执行{operation}操作",
                "filepath": filepath
            })

        except Exception as e:
            logger.error(f"PPT删除工具错误: {str(e)}")
            return json.dumps({
                "success": False,
                "message": f"错误: {str(e)}"
            })

    async def _arun(self, input_data: str) -> str:
        return self._run(input_data) 