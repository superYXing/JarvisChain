#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
PPT结构分析器 - 读取PPT文件并生成结构化表示
"""

import os
import yaml
from typing import Dict, Any, List, Optional
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from PIL import Image
import io
import base64
from src.JarvisChain.utils.logger import get_logger

logger = get_logger('ppt_structure_analyzer')

class PPTStructureAnalyzer:
    """PPT结构分析器，读取PPT并生成YAML/JSON结构表示"""
    
    def __init__(self):
        self.current_ppt_path = None
        self.current_structure = None
        
    def analyze_ppt(self, ppt_path: str) -> Dict[str, Any]:
        """分析PPT文件，返回结构化数据"""
        try:
            logger.info(f"🔍 开始分析PPT文件: {ppt_path}")
            
            if not os.path.exists(ppt_path):
                logger.error(f"❌ PPT文件不存在: {ppt_path}")
                return {"error": "PPT文件不存在"}
            
            prs = Presentation(ppt_path)
            structure = {
                "file_path": ppt_path,
                "slide_count": len(prs.slides),
                "slides": []
            }
            
            # 分析每一页幻灯片
            for idx, slide in enumerate(prs.slides):
                slide_info = self._analyze_slide(slide, idx + 1)
                structure["slides"].append(slide_info)
            
            self.current_ppt_path = ppt_path
            self.current_structure = structure
            
            logger.info(f"✅ PPT分析完成，共 {len(prs.slides)} 页")
            return structure
            
        except Exception as e:
            logger.error(f"❌ 分析PPT时发生错误: {str(e)}")
            return {"error": str(e)}
    
    def _analyze_slide(self, slide, slide_num: int) -> Dict[str, Any]:
        """分析单个幻灯片"""
        slide_info = {
            "slide_number": slide_num,
            "layout": slide.slide_layout.name if slide.slide_layout else "Unknown",
            "shapes": [],
            "has_title": False,
            "title_text": "",
            "content_summary": []
        }
        
        # 分析每个形状
        for shape in slide.shapes:
            shape_info = self._analyze_shape(shape)
            slide_info["shapes"].append(shape_info)
            
            # 提取标题
            if shape.has_text_frame and hasattr(shape, 'is_placeholder'):
                if shape.is_placeholder and shape.placeholder_format.type == 1:  # Title placeholder
                    slide_info["has_title"] = True
                    slide_info["title_text"] = shape.text_frame.text
            
            # 收集内容摘要
            if shape.has_text_frame and shape.text_frame.text.strip():
                slide_info["content_summary"].append(shape.text_frame.text.strip()[:100])
        
        return slide_info
    
    def _analyze_shape(self, shape) -> Dict[str, Any]:
        """分析单个形状"""
        shape_info = {
            "type": self._get_shape_type_name(shape),
            "name": shape.name,
            "left": shape.left,
            "top": shape.top,
            "width": shape.width,
            "height": shape.height
        }
        
        # 文本框信息
        if shape.has_text_frame:
            shape_info["has_text"] = True
            shape_info["text"] = shape.text_frame.text
            shape_info["paragraph_count"] = len(shape.text_frame.paragraphs)
        
        # 图片信息
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            shape_info["is_picture"] = True
            shape_info["image_format"] = shape.image.format if hasattr(shape.image, 'format') else "unknown"
        
        # 表格信息
        if shape.has_table:
            shape_info["has_table"] = True
            shape_info["table_rows"] = len(shape.table.rows)
            shape_info["table_columns"] = len(shape.table.columns)
        
        # 图表信息
        if shape.has_chart:
            shape_info["has_chart"] = True
            shape_info["chart_type"] = shape.chart.chart_type.name if hasattr(shape.chart, 'chart_type') else "unknown"
        
        return shape_info
    
    def _get_shape_type_name(self, shape) -> str:
        """获取形状类型名称"""
        try:
            return MSO_SHAPE_TYPE.to_xml(shape.shape_type)
        except:
            return "Unknown"
    
    def export_structure_yaml(self, output_path: Optional[str] = None) -> str:
        """导出结构为YAML格式"""
        if not self.current_structure:
            return ""
        
        yaml_content = yaml.dump(self.current_structure, 
                                allow_unicode=True, 
                                default_flow_style=False,
                                sort_keys=False)
        
        if output_path:
            with open(output_path, 'w', encoding='utf-8') as f:
                f.write(yaml_content)
            logger.info(f"📄 结构已导出到: {output_path}")
        
        return yaml_content
    
    def get_slide_preview(self, slide_number: int) -> Optional[Dict[str, Any]]:
        """获取指定幻灯片的预览信息"""
        if not self.current_structure or slide_number < 1:
            return None
        
        slides = self.current_structure.get("slides", [])
        if slide_number > len(slides):
            return None
        
        return slides[slide_number - 1]
    
    def render_slide_to_image(self, ppt_path: str, slide_number: int, output_path: Optional[str] = None) -> Optional[str]:
        """渲染幻灯片为图片（需要系统支持）"""
        try:
            # 这里可以使用python-pptx2img或其他工具
            # 暂时返回占位符
            logger.warning("⚠️ 幻灯片图片渲染功能尚未实现")
            return None
        except Exception as e:
            logger.error(f"❌ 渲染幻灯片图片失败: {str(e)}")
            return None
    
    def get_modification_suggestions(self, slide_number: int, target_changes: str) -> List[Dict[str, Any]]:
        """基于当前结构生成修改建议"""
        suggestions = []
        
        if not self.current_structure:
            return suggestions
        
        slide_info = self.get_slide_preview(slide_number)
        if not slide_info:
            return suggestions
        
        # 基于目标变更生成建议
        logger.info(f"🤔 为第 {slide_number} 页生成修改建议: {target_changes}")
        
        # 这里可以集成更智能的建议生成逻辑
        if "标题" in target_changes:
            suggestions.append({
                "type": "modify_title",
                "current": slide_info.get("title_text", ""),
                "suggestion": "修改标题文本或样式"
            })
        
        if "图片" in target_changes or "图像" in target_changes:
            picture_count = sum(1 for shape in slide_info.get("shapes", []) 
                              if shape.get("is_picture", False))
            suggestions.append({
                "type": "add_or_modify_picture",
                "current_count": picture_count,
                "suggestion": "添加或替换图片"
            })
        
        return suggestions 