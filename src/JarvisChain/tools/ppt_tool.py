#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
PPT工具 - 实现完整的PPT创建流程：大纲生成 → 代码生成 → PPT生成
"""

import json
import os
from datetime import datetime
from typing import Dict, Any, List, Optional
from langchain_openai import ChatOpenAI
from src.JarvisChain.utils.logger import get_logger
from src.JarvisChain.config.config import MODEL_CONFIG
from src.JarvisChain.agents.ppt_code_generator import PPTCodeGenerator

logger = get_logger('ppt_tool')

# 初始化PPT专用模型 - 使用analysis模型（用于内容生成）
PPT_MODEL = ChatOpenAI(
    model=MODEL_CONFIG["analysis"]["model_name"],
    temperature=0.7  # PPT创作需要一定的创意性
)

class PPTTool:
    """PPT工具：完整的PPT创建流程"""
    
    def __init__(self):
        self.name = "ppt_tool"
        self.description = "专业的PPT创建工具，支持大纲生成、代码生成和PPT文件创建"
        self.current_step = "idle"  # idle, outline, code, ppt
        self.outline_content = None
        self.code_content = None
        self.ppt_path = None
        
        # 初始化PPT代码生成器
        try:
            self.code_generator = PPTCodeGenerator()
            logger.info("✅ PPTCodeGenerator初始化成功")
        except Exception as e:
            logger.error(f"❌ PPTCodeGenerator初始化失败: {str(e)}")
            self.code_generator = None
        
    async def generate_outline(self, user_input: str, **kwargs) -> Dict[str, Any]:
        """第一步：生成PPT大纲"""
        try:
            logger.info(f"开始生成PPT大纲，用户需求: {user_input}")
            
            # 提取参数
            target_audience = kwargs.get("target_audience", "一般听众")
            slide_count = kwargs.get("slide_count", "10-15页")
            style = kwargs.get("style", "专业简洁")
            
            outline_prompt = self._build_outline_prompt(user_input, target_audience, slide_count, style)
            
            response = await PPT_MODEL.ainvoke([
                {"role": "system", "content": self._get_outline_system_prompt()},
                {"role": "user", "content": outline_prompt}
            ])
            
            self.outline_content = response.content
            self.current_step = "outline"
            
            logger.info("✅ PPT大纲生成完成")
            
            return {
                "success": True,
                "step": "outline",
                "outline": self.outline_content,
                "next_step": "生成PPT代码",
                "message": "PPT大纲已生成完成！接下来将生成PPT代码。"
            }
            
        except Exception as e:
            logger.error(f"大纲生成失败: {str(e)}")
            return {
                "success": False,
                "error": f"大纲生成失败: {str(e)}",
                "step": "outline"
            }
    
    async def generate_code(self, outline: str = None) -> Dict[str, Any]:
        """第二步：根据大纲生成PPT代码 - 使用pptCodeGenerator"""
        try:
            logger.info("开始生成PPT代码（使用PPTCodeGenerator）")
            
            # 检查代码生成器是否可用
            if not self.code_generator:
                logger.error("PPTCodeGenerator未初始化")
                return {
                    "success": False,
                    "error": "PPTCodeGenerator未初始化，请检查环境配置",
                    "step": "code"
                }
            
            # 使用当前大纲或传入的大纲
            current_outline = outline or self.outline_content
            if not current_outline:
                return {
                    "success": False,
                    "error": "没有可用的大纲，请先生成大纲",
                    "step": "code"
                }
            
            # 构建代码生成提示词
            code_prompt = f"{current_outline}"
            
            # 使用PPTCodeGenerator生成代码
            code_result = await self.code_generator.generate_ppt_code(code_prompt)
            
            if code_result["success"]:
                self.code_content = code_result["code"]
                self.current_step = "code"
                
                logger.info("✅ PPT代码生成完成（PPTCodeGenerator）")
                
                return {
                    "success": True,
                    "step": "code",
                    "code": self.code_content,
                    "next_step": "生成PPT文件",
                    "message": "PPT代码已生成完成！接下来将创建PPT文件。"
                }
            else:
                logger.error(f"PPTCodeGenerator生成失败: {code_result.get('error')}")
                return {
                    "success": False,
                    "error": f"代码生成失败: {code_result.get('error', '未知错误')}",
                    "step": "code"
                }
            
        except Exception as e:
            logger.error(f"代码生成失败: {str(e)}")
            return {
                "success": False,
                "error": f"代码生成失败: {str(e)}",
                "step": "code"
            }
    
    async def generate_ppt_file(self, code: str = None) -> Dict[str, Any]:
        """第三步：根据代码生成PPT文件 - 使用pptCodeGenerator执行"""
        try:
            logger.info("开始生成PPT文件（使用PPTCodeGenerator执行）")
            
            # 检查代码生成器是否可用
            if not self.code_generator:
                logger.error("PPTCodeGenerator未初始化")
                return {
                    "success": False,
                    "error": "PPTCodeGenerator未初始化，请检查环境配置",
                    "step": "ppt"
                }
            
            # 使用当前代码或传入的代码
            current_code = code or self.code_content
            if not current_code:
                return {
                    "success": False,
                    "error": "没有可用的PPT代码，请先生成代码",
                    "step": "ppt"
                }
            
            # 创建输出目录
            output_dir = "output/ppt"
            os.makedirs(output_dir, exist_ok=True)
            
            # 生成文件名
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            ppt_filename = f"presentation_{timestamp}.pptx"
            self.ppt_path = os.path.join(output_dir, ppt_filename)
            
            # 修改代码中的保存路径
            modified_code = self._modify_code_save_path(current_code, self.ppt_path)
            
            # 添加预置函数库到代码开头
            complete_code = self._add_prebuilt_functions_to_code(modified_code)
            
            # 使用PPTCodeGenerator执行代码
            execute_result = await self.code_generator.execute_code(complete_code)
            
            if execute_result["success"]:
                # 检查文件是否实际生成
                if os.path.exists(self.ppt_path):
                    self.current_step = "ppt"
                    logger.info(f"✅ PPT文件生成完成: {self.ppt_path}")
                    
                    return {
                        "success": True,
                        "step": "ppt",
                        "ppt_path": self.ppt_path,
                        "message": f"PPT文件已生成完成！文件路径：{self.ppt_path}"
                    }
                else:
                    logger.error(f"PPT文件未生成: {self.ppt_path}")
                    return {
                        "success": False,
                        "error": "PPT文件生成成功但文件不存在",
                        "step": "ppt"
                    }
            else:
                logger.error(f"PPTCodeGenerator执行失败: {execute_result.get('error')}")
                return {
                    "success": False,
                    "error": f"PPT文件生成失败: {execute_result.get('error', '未知错误')}",
                    "step": "ppt"
                }
                
        except Exception as e:
            logger.error(f"PPT文件生成失败: {str(e)}")
            return {
                "success": False,
                "error": f"PPT文件生成失败: {str(e)}",
                "step": "ppt"
            }
    
    def _add_prebuilt_functions_to_code(self, code: str) -> str:
        """在代码开头添加预置函数库"""
        try:
            # 检查代码是否已包含必要的导入
            has_pptx_import = "from pptx import Presentation" in code or "import pptx" in code
            has_search_function = "search_and_download_image" in code
            
            # 如果代码缺少基本导入或搜索函数，添加预置函数库
            if not has_pptx_import or not has_search_function:
                logger.info("添加预置函数库到代码")
                
                # 获取预置函数库
                if self.code_generator:
                    prebuilt_functions = self.code_generator.pptx_functions
                    
                    # 确保PPT_SAVE_DIR存在
                    if "PPT_SAVE_DIR" not in code:
                        prebuilt_functions += "\n# 确保保存目录存在\nos.makedirs(PPT_SAVE_DIR, exist_ok=True)\n"
                    
                    code = prebuilt_functions + "\n\n" + code
                else:
                    # 如果没有代码生成器，至少添加基本导入
                    logger.warning("PPTCodeGenerator不可用，添加基本导入")
                    basic_imports = """
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import os

# 创建保存目录
PPT_SAVE_DIR = "ppts"
os.makedirs(PPT_SAVE_DIR, exist_ok=True)

def search_and_download_image(query: str, filename: str, search_online=True):
    # 简化版图片搜索函数（如果没有Tavily API）
    return None

"""
                    code = basic_imports + "\n\n" + code
            else:
                logger.info("代码已包含必要的导入和函数")
            
            return code
            
        except Exception as e:
            logger.error(f"添加预置函数库失败: {str(e)}")
            # 即使失败也要添加基本导入
            basic_imports = """
from pptx import Presentation
from pptx.util import Inches, Pt
import os

PPT_SAVE_DIR = "ppts"
os.makedirs(PPT_SAVE_DIR, exist_ok=True)

"""
            return basic_imports + "\n\n" + code
    
    async def create_full_ppt(self, user_input: str, **kwargs) -> Dict[str, Any]:
        """完整流程：大纲 → 代码 → PPT"""
        try:
            logger.info("开始完整PPT创建流程")
            
            # 步骤1：生成大纲
            outline_result = await self.generate_outline(user_input, **kwargs)
            if not outline_result["success"]:
                return outline_result
            
            # 步骤2：生成代码
            code_result = await self.generate_code()
            if not code_result["success"]:
                return code_result
            
            # 步骤3：生成PPT文件
            ppt_result = await self.generate_ppt_file()
            if not ppt_result["success"]:
                return ppt_result
            
            logger.info("✅ 完整PPT创建流程完成")
            
            return {
                "success": True,
                "step": "complete",
                "outline": self.outline_content,
                "code": self.code_content,
                "ppt_path": self.ppt_path,
                "message": f"PPT创建完成！文件路径：{self.ppt_path}",
                "steps_completed": ["outline", "code", "ppt"]
            }
            
        except Exception as e:
            logger.error(f"完整PPT创建失败: {str(e)}")
            return {
                "success": False,
                "error": f"完整PPT创建失败: {str(e)}",
                "step": "failed"
            }
    
    def _build_outline_prompt(self, user_input: str, target_audience: str, slide_count: str, style: str) -> str:
        """构建大纲生成提示词"""
        return f"""用户需求: {user_input}
目标受众: {target_audience}
预期页数: {slide_count}
演示风格: {style}

请根据以上信息生成详细的PPT大纲。"""
    
    def _build_code_prompt(self, outline: str) -> str:
        """构建代码生成提示词"""
        return f"""PPT大纲:
{outline}

请根据以上大纲生成相应的PPT代码，使用python-pptx库。"""
    
    def _get_outline_system_prompt(self) -> str:
        """获取大纲生成系统提示词"""
        return """你是一个专业的PPT大纲生成专家。

请遵循以下原则生成高质量的PPT大纲：

1. **结构清晰**：使用层级化的大纲结构
2. **逻辑合理**：内容安排符合演示流程和受众心理
3. **详细具体**：每页都有明确的内容要点
4. **实用性强**：考虑实际演示效果和时间分配

大纲格式要求：
```
# PPT大纲：[主题名称]

## 基本信息
- 总页数：X页
- 建议时长：X分钟
- 目标受众：[受众描述]
- 演示风格：[风格描述]

## 详细大纲

### 第1页：封面页
- 主标题：[具体标题]
- 副标题：[副标题]
- 演示者信息
- 日期和场合

### 第2页：目录/议程
- 主要章节列表
- 预计时间分配
- 重点内容预告

### 第3页：[章节名称]
- 核心观点1
- 核心观点2
- 支撑数据或案例
- 视觉元素建议

[继续其他页面...]

## 演示建议
- 重点强调部分
- 互动环节设计
- 可能的问答准备
```

请确保大纲内容丰富、结构合理、易于实现。"""
    
    def _get_code_system_prompt(self) -> str:
        """获取代码生成系统提示词"""
        return """你是一个专业的PPT代码生成专家，精通python-pptx库。

请根据提供的PPT大纲生成完整的Python代码，要求：

1. **完整可运行**：生成的代码可以直接执行
2. **结构清晰**：代码组织良好，有适当的注释
3. **样式美观**：使用合适的字体、颜色、布局
4. **内容丰富**：实现大纲中的所有要点

代码结构要求：
```python
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

def create_presentation():
    # 创建演示文稿
    prs = Presentation()
    
    # 设置幻灯片尺寸（16:9）
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    
    # 第1页：封面页
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # 空白布局
    # 添加标题、副标题等
    
    # 第2页：目录页
    slide = prs.slides.add_slide(prs.slide_layouts[1])  # 标题和内容布局
    # 添加目录内容
    
    # [继续其他页面...]
    
    return prs

# 主函数
if __name__ == "__main__":
    prs = create_presentation()
    prs.save("presentation.pptx")
    print("PPT生成完成！")
```

重要提示：
- 使用Inches()设置尺寸
- 使用Pt()设置字体大小
- 合理使用颜色和对齐方式
- 每页都要有明确的主题和内容
- 代码必须完整且可执行"""
    
    def _modify_code_save_path(self, code: str, output_path: str) -> str:
        """修改代码中的保存路径"""
        try:
            # 清理代码（移除markdown格式）
            if code.startswith('```python'):
                code = code[9:]
            if code.endswith('```'):
                code = code[:-3]
            code = code.strip()
            
            # 替换保存路径 - 修复Windows路径转义问题
            # 使用原始字符串避免转义问题
            import re
            
            # 将Windows路径转换为正斜杠格式
            safe_path = output_path.replace('\\', '/')
            
            # 匹配 prs.save("filename") 模式
            code = re.sub(r'prs\.save\(["\'][^"\']*["\']\)', f'prs.save("{safe_path}")', code)
            
            # 匹配 save_presentation(filename) 模式  
            code = re.sub(r'save_presentation\(["\'][^"\']*["\']\)', f'save_presentation("{safe_path}")', code)
            
            # 匹配PPT_SAVE_DIR相关的保存语句
            code = re.sub(r'prs\.save\(os\.path\.join\(PPT_SAVE_DIR[^)]*\)\)', f'prs.save("{safe_path}")', code)
            
            # 如果没有找到保存语句，在代码末尾添加
            if 'prs.save(' not in code and 'save_presentation(' not in code:
                if 'prs =' in code or 'prs=' in code:
                    code += f'\nprs.save("{safe_path}")\nprint("PPT已保存到: {safe_path}")'
            
            return code
            
        except Exception as e:
            logger.error(f"修改代码保存路径失败: {str(e)}")
            return code

    async def _create_ppt_file(self, code: str, output_path: str) -> bool:
        """执行PPT代码生成文件（已弃用，现在使用PPTCodeGenerator）"""
        logger.warning("_create_ppt_file方法已弃用，现在使用PPTCodeGenerator.execute_code")
        return False
    
    def get_current_step(self) -> str:
        """获取当前步骤"""
        return self.current_step
    
    def get_outline(self) -> Optional[str]:
        """获取当前大纲"""
        return self.outline_content
    
    def get_code(self) -> Optional[str]:
        """获取当前代码"""
        return self.code_content
    
    def get_ppt_path(self) -> Optional[str]:
        """获取PPT文件路径"""
        return self.ppt_path
    
    def reset(self):
        """重置工具状态"""
        self.current_step = "idle"
        self.outline_content = None
        self.code_content = None
        self.ppt_path = None