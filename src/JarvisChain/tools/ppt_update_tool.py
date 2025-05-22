from langchain.tools import BaseTool
from pydantic import BaseModel, Field
import json
from src.JarvisChain.utils.logger import get_logger
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
import os

# Get logger
logger = get_logger('ppt_update')



class PPTUpdateTool(BaseTool):
    name: str = "ppt_update_tool"
    description: str = """用于更新PPT元素的工具，包括修改文本、图片、背景等。"""
  
    ppts_dir: str = os.path.join(os.getcwd(), "ppts")

    def __init__(self, **data):
        super().__init__(**data)
        if not os.path.exists(self.ppts_dir):
            os.makedirs(self.ppts_dir)

    def _load_presentation(self, filename: str) -> Presentation:
        """加载演示文稿"""
        filepath = os.path.join(self.ppts_dir, filename)
        if not os.path.exists(filepath):
            raise FileNotFoundError(f"未找到演示文稿: {filename}")
        return Presentation(filepath)

    def _save_presentation(self, prs: Presentation, filename: str) -> str:
        """保存演示文稿"""
        filepath = os.path.join(self.ppts_dir, filename)
        prs.save(filepath)
        return filepath

    def _update_text(self, slide, shape_index: int, text: str, style: dict = None) -> None:
        """更新文本内容"""
        if 0 <= shape_index < len(slide.shapes):
            shape = slide.shapes[shape_index]
            if hasattr(shape, "text_frame"):
                text_frame = shape.text_frame
                text_frame.text = text
                
                if style:
                    paragraph = text_frame.paragraphs[0]
                    # 设置对齐方式
                    if 'align' in style:
                        paragraph.alignment = PP_ALIGN.CENTER if style['align'] == 'center' else PP_ALIGN.LEFT
                    
                    # 设置字体
                    if 'font' in style or 'size' in style or 'color' in style:
                        font = paragraph.font
                        if 'font' in style:
                            font.name = style['font']
                        if 'size' in style:
                            font.size = Pt(style['size'])
                        if 'color' in style:
                            r, g, b = style['color']
                            font.color.rgb = RGBColor(r, g, b)

    def _update_image(self, slide, shape_index: int, image_path: str, position: dict = None) -> None:
        """更新图片"""
        if 0 <= shape_index < len(slide.shapes):
            shape = slide.shapes[shape_index]
            if shape.shape_type == 13:  # 图片类型
                # 删除旧图片
                sp = shape._element
                sp.getparent().remove(sp)
                
                # 添加新图片
                if position:
                    left = Inches(position.get('left', 1))
                    top = Inches(position.get('top', 1))
                    width = Inches(position.get('width', 6))
                    height = Inches(position.get('height', 4))
                else:
                    left = Inches(1)
                    top = Inches(1)
                    width = Inches(6)
                    height = Inches(4)
                
                slide.shapes.add_picture(image_path, left, top, width, height)

    def _update_background(self, slide, color: tuple) -> None:
        """更新背景颜色"""
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(*color)

    def _run(self, input_data: str) -> str:
        """处理更新PPT元素的命令"""
        try:
            # 解析 JSON 查询
            data = json.loads(input_data)
            operation = data.get("operation")
            parameters = data.get("parameters", {})
            
            # 执行操作
            prs = self._load_presentation(parameters.get("filename"))
            slide_index = parameters.get("slide_index", 0)
            
            if 0 <= slide_index < len(prs.slides):
                slide = prs.slides[slide_index]
                
                if operation == "update_text":
                    shape_index = parameters.get("shape_index", 0)
                    self._update_text(slide, shape_index, 
                                    parameters.get("content", ""),
                                    parameters.get("style"))
                elif operation == "update_image":
                    shape_index = parameters.get("shape_index", 0)
                    self._update_image(slide, shape_index,
                                     parameters.get("image_path"),
                                     parameters.get("position"))
                elif operation == "update_background":
                    self._update_background(slide, parameters.get("color", (255, 255, 255)))
            
            # 保存演示文稿
            filepath = self._save_presentation(prs, parameters.get("filename"))
            
            return json.dumps({
                "success": True,
                "operation": operation,
                "message": f"成功执行{operation}操作",
                "filepath": filepath
            })

        except Exception as e:
            logger.error(f"PPT更新工具错误: {str(e)}")
            return json.dumps({
                "success": False,
                "message": f"错误: {str(e)}"
            })

    async def _arun(self, input_data: str) -> str:
        return self._run(input_data) 